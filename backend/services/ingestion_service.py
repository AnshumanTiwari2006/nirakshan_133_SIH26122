import io
import logging
from typing import Optional, Dict, Any
import docx
import openpyxl
import pypdf

logger = logging.getLogger(__name__)

class IngestionService:
    def parse_txt(self, content: bytes) -> str:
        return content.decode("utf-8", errors="ignore")
    
    def parse_docx(self, content: bytes) -> str:
        doc = docx.Document(io.BytesIO(content))
        return "\n".join([para.text for para in doc.paragraphs])
    
    def parse_xlsx(self, content: bytes) -> str:
        wb = openpyxl.load_workbook(io.BytesIO(content))
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) for cell in row if cell is not None])
                if row_text:
                    rows.append(row_text)
        return "\n".join(rows)
    
    def parse_pdf(self, content: bytes) -> str:
        reader = pypdf.PdfReader(io.BytesIO(content))
        return "\n".join([page.extract_text() or "" for page in reader.pages])
    
    def parse(self, content: bytes, file_type: str) -> str:
        parsers = {
            "TXT": self.parse_txt,
            "DOCX": self.parse_docx,
            "XLSX": self.parse_xlsx,
            "PDF": self.parse_pdf,
        }
        parser = parsers.get(file_type.upper())
        if not parser:
            raise ValueError(f"Unsupported file type: {file_type}")
        return parser(content)

ingestion_service = IngestionService()